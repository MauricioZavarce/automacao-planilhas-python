import argparse
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------
# Funções auxiliares
# ------------------------------------------------------------

def ler_planilha(caminho: Path) -> pd.DataFrame:
    """Lê uma planilha Excel ou CSV e retorna um DataFrame."""
    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")
    if caminho.suffix.lower() in ['.xlsx', '.xls']:
        return pd.read_excel(caminho)
    elif caminho.suffix.lower() == '.csv':
        return pd.read_csv(caminho)
    else:
        raise ValueError(f"Formato de arquivo não suportado: {caminho.suffix}")

def processar_dados(df: pd.DataFrame) -> pd.DataFrame:
    """Aplica transformações: aqui calcula Total = Preco * Quantidade."""
    if 'Preco' not in df.columns or 'Quantidade' not in df.columns:
        raise ValueError("A planilha deve conter as colunas 'Preco' e 'Quantidade'.")
    df['Total'] = df['Preco'] * df['Quantidade']
    return df

def salvar_planilha(df: pd.DataFrame, caminho: Path, formatar: bool = True) -> None:
    """Salva o DataFrame em Excel e aplica formatação opcional."""
    # Salva temporariamente sem formatação
    df.to_excel(caminho, index=False)
    if not formatar:
        return

    # Carrega para formatação com openpyxl
    wb = load_workbook(caminho)
    ws = wb.active

    # Negrito no cabeçalho
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # Destacar células da coluna 'Total' maiores que um limite (ex.: 10000)
    col_idx_total = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == 'Total':
            col_idx_total = idx
            break

    if col_idx_total:
        for row in ws.iter_rows(min_row=2, min_col=col_idx_total, max_col=col_idx_total):
            for cell in row:
                if cell.value and cell.value > 10000:
                    cell.fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    # Ajustar largura das colunas
    for col in ws.columns:
        max_length = 0
        column_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column_letter].width = max_length + 2

    wb.save(caminho)

# ------------------------------------------------------------
# Função principal com argparse
# ------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description='Automatiza processamento de planilhas: calcula Total e formata o resultado.'
    )
    parser.add_argument(
        'entrada',
        type=str,
        help='Caminho do arquivo de entrada (Excel ou CSV)'
    )
    parser.add_argument(
        '-o', '--saida',
        type=str,
        default=None,
        help='Caminho do arquivo de saída (padrão: entrada_processada.xlsx)'
    )
    parser.add_argument(
        '--sem-formatacao',
        action='store_true',
        help='Não aplica formatação (negrito, cores, largura)'
    )
    args = parser.parse_args()

    caminho_entrada = Path(args.entrada)
    if args.saida:
        caminho_saida = Path(args.saida)
    else:
        caminho_saida = caminho_entrada.with_name(caminho_entrada.stem + '_processada.xlsx')

    try:
        print(f"Lendo arquivo: {caminho_entrada}")
        df = ler_planilha(caminho_entrada)

        print("Processando dados...")
        df = processar_dados(df)

        print(f"Salvando em: {caminho_saida}")
        salvar_planilha(df, caminho_saida, formatar=not args.sem_formatacao)

        print("✅ Concluído com sucesso!")
    except Exception as e:
        print(f"❌ Erro: {e}", file=sys.stderr)
        sys.exit(1)

if __name__ == '__main__':
    main()